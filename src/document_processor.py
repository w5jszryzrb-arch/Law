import re
import hashlib
from pathlib import Path
from dataclasses import dataclass, field
from typing import Optional

import fitz  # PyMuPDF
from pptx import Presentation

import config


@dataclass
class DocumentChunk:
    text: str
    chunk_index: int
    doc_id: str
    source_filename: str
    doc_type: str
    title: str
    section: str = ""
    metadata: dict = field(default_factory=dict)


@dataclass
class ProcessedDocument:
    doc_id: str
    title: str
    doc_type: str
    source_filename: str
    full_text: str
    chunks: list[DocumentChunk]
    page_count: int = 0
    metadata: dict = field(default_factory=dict)


def make_doc_id(filename: str, title: str) -> str:
    return hashlib.md5(f"{filename}:{title}".encode()).hexdigest()[:16]


def extract_pdf(path: str) -> tuple[str, int]:
    doc = fitz.open(path)
    pages = []
    for i, page in enumerate(doc):
        text = page.get_text("text")
        text = re.sub(r"\n{3,}", "\n\n", text)
        if text.strip():
            pages.append(text)
    return "\n\n".join(pages), len(doc)


def extract_pptx(path: str) -> tuple[str, int]:
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        body_parts = []
        notes_text = ""

        if slide.shapes.title and slide.shapes.title.text.strip():
            title = slide.shapes.title.text.strip()

        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        body_parts.append(t)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                notes_text = f"\n[Notes: {notes}]"

        slide_text = f"[Slide {i}: {title}]\n" + "\n".join(body_parts) + notes_text
        slides.append(slide_text)

    return "\n\n".join(slides), len(prs.slides)


def extract_txt(path: str) -> tuple[str, int]:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return text, 0


def extract_text(path: str) -> tuple[str, int]:
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    elif suffix in (".pptx", ".ppt"):
        return extract_pptx(path)
    else:
        return extract_txt(path)


def _detect_section(text: str) -> str:
    patterns = [
        r"^\[Slide \d+[:\]]",
        r"^(?:Chapter|Section|Part)\s+\w+",
        r"^\d+\.\s+[A-Z]",
    ]
    for p in patterns:
        m = re.match(p, text.strip(), re.IGNORECASE)
        if m:
            return m.group(0)[:80]
    return ""


def chunk_text(
    text: str,
    doc_id: str,
    source_filename: str,
    doc_type: str,
    title: str,
    chunk_size: int = config.CHUNK_SIZE,
    overlap: int = config.CHUNK_OVERLAP,
) -> list[DocumentChunk]:
    paragraphs = re.split(r"\n\n+", text)
    chunks: list[DocumentChunk] = []
    current: list[str] = []
    current_len = 0

    def flush(idx: int) -> None:
        combined = "\n\n".join(current).strip()
        if not combined:
            return
        chunks.append(
            DocumentChunk(
                text=combined,
                chunk_index=idx,
                doc_id=doc_id,
                source_filename=source_filename,
                doc_type=doc_type,
                title=title,
                section=_detect_section(combined),
            )
        )

    chunk_idx = 0
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        words = para.split()
        para_len = len(words)

        if para_len > chunk_size:
            # Flush current before handling an oversized paragraph
            if current:
                flush(chunk_idx)
                chunk_idx += 1
                current = current[-overlap:] if overlap else []
                current_len = sum(len(c.split()) for c in current)

            # Split the oversized paragraph into sub-chunks
            sub_words = words
            while sub_words:
                sub = " ".join(sub_words[:chunk_size])
                chunks.append(
                    DocumentChunk(
                        text=sub,
                        chunk_index=chunk_idx,
                        doc_id=doc_id,
                        source_filename=source_filename,
                        doc_type=doc_type,
                        title=title,
                        section=_detect_section(sub),
                    )
                )
                chunk_idx += 1
                sub_words = sub_words[chunk_size - overlap :]
        else:
            if current_len + para_len > chunk_size and current:
                flush(chunk_idx)
                chunk_idx += 1
                current = current[-1:] if current else []
                current_len = sum(len(c.split()) for c in current)
            current.append(para)
            current_len += para_len

    if current:
        flush(chunk_idx)

    return chunks


def auto_detect_doc_type(filename: str, text: str) -> str:
    name = filename.lower()
    sample = text[:2000].lower()

    if any(k in name for k in ("past paper", "exam", "examination", "test")):
        return "past_paper"
    if any(k in sample for k in ("question 1", "answer all questions", "examination paper")):
        return "past_paper"
    if any(k in name for k in ("lecture", "slide", "week", "tutorial")):
        return "lecture"
    if any(k in name for k in ("act ", "act_", "statute", "legislation", "regulation")):
        return "statute"
    if any(k in sample for k in ("[nzsc", "[nzca", "[nzhc", "v ", "ratio decidendi", "obiter")):
        return "case_law"
    if any(k in name for k in ("instructions", "assignment", "problem", "scenario")):
        return "instruction"
    if any(k in name for k in ("journal", "article", "review", "scholar")):
        return "article"
    return "other"


def process_uploaded_file(
    file_bytes: bytes,
    original_filename: str,
    title: str,
    doc_type: Optional[str] = None,
) -> ProcessedDocument:
    suffix = Path(original_filename).suffix.lower()
    save_path = config.UPLOAD_DIR / original_filename
    save_path.write_bytes(file_bytes)

    full_text, page_count = extract_text(str(save_path))

    if not doc_type:
        doc_type = auto_detect_doc_type(original_filename, full_text)

    doc_id = make_doc_id(original_filename, title)

    chunks = chunk_text(
        text=full_text,
        doc_id=doc_id,
        source_filename=original_filename,
        doc_type=doc_type,
        title=title,
    )

    return ProcessedDocument(
        doc_id=doc_id,
        title=title,
        doc_type=doc_type,
        source_filename=original_filename,
        full_text=full_text,
        chunks=chunks,
        page_count=page_count,
        metadata={"suffix": suffix},
    )
